import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def init_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return doc

def add_header_banner(doc, report_type):
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_before = Pt(0)
    p_uni.paragraph_format.space_after = Pt(2)
    r_uni = p_uni.add_run("CHANDIGARH UNIVERSITY")
    r_uni.bold = True
    r_uni.font.size = Pt(20)
    r_uni.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run("Discover. Learn. Empower.")
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    p_rep = doc.add_paragraph()
    p_rep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rep.paragraph_format.space_after = Pt(10)
    r_rep = p_rep.add_run(report_type)
    r_rep.bold = True
    r_rep.font.size = Pt(13)
    r_rep.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_course.paragraph_format.space_after = Pt(12)
    r_c = p_course.add_run("Course: Python Programming & Data Science — [Course Code: 25CAT-204]")
    r_c.font.size = Pt(11)
    r_c.font.italic = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(10)
    r_t = p_title.add_run("“STUDENT PERFORMANCE DATA ANALYSIS: EXPLORATORY DATA ANALYSIS, STATISTICAL INFERENCE & MULTIVARIATE VISUALIZATION”")
    r_t.bold = True
    r_t.font.size = Pt(13.5)
    r_t.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

def add_student_table(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=5, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    data = [
        ("SUBMITTED BY (STUDENT DETAILS)", True, "1B365D", "FFFFFF"),
        ("Student Name:  Mithun", False, "FFFFFF", "000000"),
        ("UID: 25BCD10179", False, "FFFFFF", "000000"),
        ("Sem: 3 |  Sec: 25BCD-2-B", False, "FFFFFF", "000000"),
        ("Branch / Program:  BCA (DATA SCIENCE)", False, "FFFFFF", "000000")
    ]

    for i, (text, is_hdr, bg, fg) in enumerate(data):
        cell = table.cell(i, 0)
        cell.width = Inches(4.5)
        set_cell_background(cell, bg)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(4)
        run = cp.add_run(text)
        run.bold = is_hdr
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor.from_string(fg)

    p_ay = doc.add_paragraph()
    p_ay.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ay.paragraph_format.space_before = Pt(24)
    r_ay = p_ay.add_run("ACADEMIC YEAR: 2026 – 2027")
    r_ay.bold = True
    r_ay.font.size = Pt(11)

    doc.add_page_break()

def add_styled_table(doc, headers, rows_data):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True

    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        set_cell_background(cell, "1B365D")
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(3)
        r = cp.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row in enumerate(rows_data):
        bg = "F4F7FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            r = cp.add_run(str(val))
            r.font.size = Pt(9.0)

    doc.add_paragraph()

def add_code_block(doc, code_str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

def add_h1(doc, text):
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    run = h.add_run(text)
    run.font.size = Pt(16.5)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

def add_h2(doc, text):
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(13)
    h.paragraph_format.space_after = Pt(4)
    run = h.add_run(text)
    run.font.size = Pt(13.0)
    run.bold = True
    run.font.color.rgb = RGBColor(0x2B, 0x4C, 0x7E)

def add_h3(doc, text):
    h = doc.add_heading(level=3)
    h.paragraph_format.space_before = Pt(9)
    h.paragraph_format.space_after = Pt(3)
    run = h.add_run(text)
    run.font.size = Pt(11.0)
    run.bold = True
    run.font.color.rgb = RGBColor(0x3B, 0x62, 0x9B)

def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)
    return p

def add_figure(doc, img_name, caption):
    base_dir = r"C:\Users\Rahul\.gemini\antigravity\scratch\sem-project\python-student-performance\figures"
    img_path = os.path.join(base_dir, img_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        doc.add_picture(img_path, width=Inches(5.3))
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(10)
        r = cp.add_run(f"Figure: {caption}")
        r.font.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
